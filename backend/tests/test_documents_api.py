"""Test API nạp tài liệu bằng store/embedder giả: không cần Postgres hay endpoint embedding."""

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.documents.router import router
from app.documents.text_extraction import split_into_chunks
from app.embedding_client import EmbeddingError
from app.settings import get_settings

from settings_factory import make_settings


class FakeStore:
    def __init__(self):
        self.docs: dict[str, tuple[str, list[str], list[list[float]]]] = {}
        self.languages: dict[str, str | None] = {}

    def get_content_hash(self, source):
        return self.docs[source][0] if source in self.docs else None

    def replace_document(self, source, language, content_hash, chunks, vectors):
        self.docs[source] = (content_hash, chunks, vectors)
        self.languages[source] = language

    def delete_document(self, source):
        return self.docs.pop(source, None) is not None


class FakeEmbedder:
    def __init__(self):
        self.calls = 0

    def embed(self, texts):
        self.calls += 1
        return [[0.0] * 1024 for _ in texts]


class FailingEmbedder:
    def embed(self, texts):
        raise EmbeddingError("endpoint down")


def make_client(tmp_path, embedder=None, **overrides):
    settings = make_settings(documents_dir=str(tmp_path), **overrides)
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_settings] = lambda: settings
    store, embedder = FakeStore(), embedder or FakeEmbedder()
    app.state.document_store, app.state.embedder = store, embedder
    return TestClient(app), store, embedder


def test_upload_indexes_text_file(tmp_path):
    client, store, _ = make_client(tmp_path)
    r = client.post("/documents/upload", files=[("files", ("faq.md", b"# Doi tra\n\nDoi tra trong 7 ngay.", "text/markdown"))])
    assert r.status_code == 200
    assert r.json()["results"] == [{"source": "faq.md", "status": "indexed", "chunks": 1, "detail": None}]
    assert "faq.md" in store.docs


def test_reupload_same_content_is_unchanged_and_skips_embedding(tmp_path):
    client, _, embedder = make_client(tmp_path)
    files = [("files", ("a.txt", b"noi dung", "text/plain"))]
    client.post("/documents/upload", files=files)
    r = client.post("/documents/upload", files=files)
    assert r.json()["results"][0]["status"] == "unchanged"
    assert embedder.calls == 1


def test_upload_strips_client_path_and_rejects_bad_types(tmp_path):
    client, store, _ = make_client(tmp_path)
    r = client.post("/documents/upload", files=[
        ("files", ("../../etc/notes.txt", b"hello", "text/plain")),
        ("files", ("virus.exe", b"MZ", "application/octet-stream")),
    ])
    results = r.json()["results"]
    assert results[0]["source"] == "notes.txt" and results[0]["status"] == "indexed"
    assert results[1]["status"] == "skipped"
    assert list(store.docs) == ["notes.txt"]


def test_upload_over_size_limit_is_skipped(tmp_path):
    client, store, _ = make_client(tmp_path, max_upload_mb=1)
    r = client.post("/documents/upload", files=[("files", ("big.txt", b"a" * (1024 * 1024 + 1), "text/plain"))])
    assert r.json()["results"][0]["status"] == "skipped"
    assert not store.docs


def test_embedding_failure_is_reported_per_file(tmp_path):
    client, store, _ = make_client(tmp_path, embedder=FailingEmbedder())
    r = client.post("/documents/upload", files=[("files", ("a.txt", b"x", "text/plain"))])
    assert r.status_code == 200
    assert r.json()["results"][0]["status"] == "failed"
    assert not store.docs


def test_non_utf8_and_corrupt_pdf_are_skipped(tmp_path):
    client, _, _ = make_client(tmp_path)
    r = client.post("/documents/upload", files=[
        ("files", ("latin.txt", "é".encode("latin-1"), "text/plain")),
        ("files", ("broken.pdf", b"not a pdf", "application/pdf")),
    ])
    assert [x["status"] for x in r.json()["results"]] == ["skipped", "skipped"]


def test_api_key_required_when_configured(tmp_path):
    client, _, _ = make_client(tmp_path, api_key="secret")
    files = [("files", ("a.txt", b"x", "text/plain"))]
    assert client.post("/documents/upload", files=files).status_code == 401
    assert client.post("/documents/upload", files=files, headers={"X-API-Key": "wrong"}).status_code == 401
    assert client.post("/documents/upload", files=files, headers={"X-API-Key": "secret"}).status_code == 200


def test_ingest_folder_recursive_and_relative_sources(tmp_path):
    (tmp_path / "sub").mkdir()
    (tmp_path / "a.txt").write_text("alpha", encoding="utf-8")
    (tmp_path / "sub" / "b.md").write_text("beta", encoding="utf-8")
    (tmp_path / "ignore.png").write_bytes(b"png")
    client, store, _ = make_client(tmp_path)
    r = client.post("/documents/ingest-folder")  # không body: dùng mặc định
    assert {x["source"] for x in r.json()["results"]} == {"a.txt", "sub/b.md"}
    assert set(store.docs) == {"a.txt", "sub/b.md"}
    flat = client.post("/documents/ingest-folder", json={"recursive": False})
    assert [x["source"] for x in flat.json()["results"]] == ["a.txt"]


@pytest.mark.parametrize("subfolder", ["..", "../outside", "/etc", "sub/../.."])
def test_ingest_folder_rejects_path_traversal(tmp_path, subfolder):
    docs = tmp_path / "docs"
    docs.mkdir()
    client, _, _ = make_client(docs)
    r = client.post("/documents/ingest-folder", json={"subfolder": subfolder})
    assert r.status_code in (400, 404)
    assert r.status_code != 200


def test_ingest_folder_missing_subfolder_is_404(tmp_path):
    client, _, _ = make_client(tmp_path)
    assert client.post("/documents/ingest-folder", json={"subfolder": "nope"}).status_code == 404


def test_concurrency_limit_returns_429(tmp_path):
    from app.documents.router import _upload_gate

    _upload_gate.cache_clear()
    client, _, _ = make_client(tmp_path, max_concurrent_uploads=1)
    gate = _upload_gate(1)
    gate.acquire()
    try:
        r = client.post("/documents/upload", files=[("files", ("a.txt", b"x", "text/plain"))])
        assert r.status_code == 429
    finally:
        gate.release()
        _upload_gate.cache_clear()


def test_split_into_chunks_respects_size_and_drops_blanks():
    text = "\n\n".join(f"Doan {i}. " + "tu " * 100 for i in range(10))
    chunks = split_into_chunks(text, chunk_size=400, chunk_overlap=50)
    assert len(chunks) > 1
    assert all(0 < len(c) <= 400 for c in chunks)
    assert split_into_chunks("   \n\n  ", 400, 50) == []


def _docx_bytes() -> bytes:
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_heading("Chinh sach doi tra", level=1)
    doc.add_paragraph("Doi tra trong 7 ngay.")
    doc.add_paragraph("San pham loi", style="List Bullet")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Loai", "Phi"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_is_converted_to_markdown_and_indexed_with_language(tmp_path):
    from app.documents.text_extraction import extract_text

    text = extract_text("a.docx", _docx_bytes())
    assert text.startswith("# Chinh sach doi tra") and "- San pham loi" in text and "| Loai | Phi |" in text
    client, store, _ = make_client(tmp_path)
    files = [("files", ("a.docx", _docx_bytes(), "application/octet-stream"))]
    r = client.post("/documents/upload", data={"language": "ko"}, files=files)
    assert r.json()["results"][0]["status"] == "indexed"
    assert store.languages["a.docx"] == "ko"


def test_invalid_language_is_rejected_and_language_change_reindexes(tmp_path):
    client, _, embedder = make_client(tmp_path)
    files = [("files", ("a.txt", b"noi dung", "text/plain"))]
    assert client.post("/documents/upload", data={"language": "fr"}, files=files).status_code == 422
    client.post("/documents/upload", data={"language": "vi"}, files=files)
    r = client.post("/documents/upload", data={"language": "ko"}, files=files)
    assert r.json()["results"][0]["status"] == "indexed" and embedder.calls == 2


def test_markdown_chunks_split_by_heading_and_keep_breadcrumb():
    text = "# Hoan tien\n\n" + "Noi dung. " * 80 + "\n\n## Dieu kien\n\n" + "Chi tiet. " * 80
    chunks = split_into_chunks(text, chunk_size=300, chunk_overlap=0, markdown=True)
    assert all(len(c) <= 300 + 40 for c in chunks)
    assert not any("Noi dung" in c and "Chi tiet" in c for c in chunks)  # không vắt qua hai mục
    assert any(c.startswith("Hoan tien > Dieu kien") for c in chunks)


def test_delete_document(tmp_path):
    client, store, _ = make_client(tmp_path)
    client.post("/documents/upload", files=[("files", ("a.txt", b"x", "text/plain"))])
    assert client.delete("/documents", params={"source": "a.txt"}).status_code == 200
    assert not store.docs
    assert client.delete("/documents", params={"source": "a.txt"}).status_code == 404


def test_embedder_accepts_base_and_full_url():
    from app.embedding_client import HttpEmbedder

    for url in ("http://h/v1", "http://h/v1/", "http://h/v1/embeddings"):
        assert HttpEmbedder(url, "m", 1024, 10, 5)._url == "http://h/v1/embeddings"
